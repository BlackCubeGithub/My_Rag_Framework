"""
Document Parser
Multi-format document parsing (PDF, Word, Markdown, HTML)
"""
import re
import structlog
from pathlib import Path
from typing import Optional
from dataclasses import dataclass, field

logger = structlog.get_logger()


@dataclass
class ParsedDocument:
    """Container for parsed document content"""
    title: str = ""
    content: str = ""
    text_blocks: list[dict] = field(default_factory=list)
    tables: list[dict] = field(default_factory=list)
    images: list[dict] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


class DocumentParser:
    """Parses various document formats"""

    SUPPORTED_FORMATS = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".doc": "docx",
        ".txt": "text",
        ".md": "markdown",
        ".html": "html",
        ".htm": "html",
    }

    def __init__(self):
        self.logger = logger

    async def parse(self, file_path: str, file_name: str) -> ParsedDocument:
        """
        Parse a document based on its file type.

        Args:
            file_path: Path to the document
            file_name: Original file name

        Returns:
            ParsedDocument with extracted content
        """
        path = Path(file_path)
        suffix = path.suffix.lower()

        parser_type = self.SUPPORTED_FORMATS.get(suffix)
        if not parser_type:
            raise ValueError(f"Unsupported file format: {suffix}")

        self.logger.info("parsing_document", file_name=file_name, type=parser_type)

        if parser_type == "pdf":
            return await self._parse_pdf(file_path, file_name)
        elif parser_type == "docx":
            return await self._parse_docx(file_path, file_name)
        elif parser_type == "markdown":
            return await self._parse_markdown(file_path, file_name)
        elif parser_type == "html":
            return await self._parse_html(file_path, file_name)
        elif parser_type == "text":
            return await self._parse_text(file_path, file_name)

        return ParsedDocument(metadata={"file_name": file_name})

    async def _parse_pdf(self, file_path: str, file_name: str) -> ParsedDocument:
        """Parse PDF document with text and structure extraction"""
        doc = ParsedDocument()
        doc.metadata["file_name"] = file_name
        doc.metadata["source"] = "pdf"

        try:
            import pdfplumber

            with pdfplumber.open(file_path) as pdf:
                doc.metadata["page_count"] = len(pdf.pages)

                for page_num, page in enumerate(pdf.pages, 1):
                    page_text = page.extract_text() or ""

                    blocks = self._split_into_blocks(page_text)
                    for block in blocks:
                        doc.text_blocks.append({
                            "text": block,
                            "page": page_num,
                            "type": "text",
                        })

                    doc.content += page_text + "\n\n"

                    tables = page.extract_tables()
                    for table in tables:
                        if table:
                            doc.tables.append({
                                "data": table,
                                "page": page_num,
                            })

        except ImportError:
            self.logger.warning("pdfplumber not available, using fallback")
            return await self._parse_pdf_fallback(file_path, file_name)
        except Exception as e:
            self.logger.error("pdf_parsing_failed", error=str(e))
            return await self._parse_pdf_fallback(file_path, file_name)

        doc.title = self._extract_title(doc.content) or file_name
        return doc

    async def _parse_pdf_fallback(self, file_path: str, file_name: str) -> ParsedDocument:
        """Fallback PDF parsing using PyMuPDF"""
        doc = ParsedDocument()
        doc.metadata["file_name"] = file_name

        try:
            import fitz

            pdf_doc = fitz.open(file_path)
            doc.metadata["page_count"] = len(pdf_doc)

            for page_num in range(len(pdf_doc)):
                page = pdf_doc[page_num]
                text = page.get_text()
                doc.content += text + "\n\n"

                doc.text_blocks.append({
                    "text": text,
                    "page": page_num + 1,
                    "type": "text",
                })

        except ImportError:
            self.logger.error("Neither pdfplumber nor PyMuPDF available")
        except Exception as e:
            self.logger.error("fallback_pdf_parsing_failed", error=str(e))

        doc.title = file_name
        return doc

    async def _parse_docx(self, file_path: str, file_name: str) -> ParsedDocument:
        """Parse Word document"""
        doc = ParsedDocument()
        doc.metadata["file_name"] = file_name

        try:
            from docx import Document

            docx_doc = Document(file_path)

            doc.title = docx_doc.core_properties.title or file_name

            for para in docx_doc.paragraphs:
                text = para.text.strip()
                if text:
                    doc.text_blocks.append({
                        "text": text,
                        "type": "text",
                        "style": para.style.name,
                    })
                    doc.content += text + "\n"

            for table in docx_doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)

                if table_data:
                    doc.tables.append({"data": table_data})

        except ImportError:
            self.logger.error("python-docx not available")
        except Exception as e:
            self.logger.error("docx_parsing_failed", error=str(e))

        return doc

    async def _parse_markdown(self, file_path: str, file_name: str) -> ParsedDocument:
        """Parse Markdown document"""
        doc = ParsedDocument()
        doc.metadata["file_name"] = file_name

        try:
            content = Path(file_path).read_text(encoding="utf-8")

            lines = content.split("\n")
            in_code_block = False
            current_block = []

            for i, line in enumerate(lines, 1):
                if line.startswith("```"):
                    in_code_block = not in_code_block
                    continue

                if line.startswith("#"):
                    doc.title = doc.title or line.lstrip("#").strip()
                    if current_block:
                        doc.text_blocks.append({
                            "text": "\n".join(current_block),
                            "type": "text",
                        })
                        current_block = []
                else:
                    current_block.append(line)

                if not in_code_block:
                    doc.content += line + "\n"

            if current_block:
                doc.text_blocks.append({
                    "text": "\n".join(current_block),
                    "type": "text",
                })

        except Exception as e:
            self.logger.error("markdown_parsing_failed", error=str(e))

        if not doc.title:
            doc.title = file_name

        return doc

    async def _parse_html(self, file_path: str, file_name: str) -> ParsedDocument:
        """Parse HTML document"""
        doc = ParsedDocument()
        doc.metadata["file_name"] = file_name

        try:
            from bs4 import BeautifulSoup

            with open(file_path, "r", encoding="utf-8") as f:
                html_content = f.read()

            soup = BeautifulSoup(html_content, "html.parser")

            doc.title = soup.title.string if soup.title else file_name

            for tag in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "td", "th"]):
                text = tag.get_text().strip()
                if text:
                    doc.text_blocks.append({
                        "text": text,
                        "type": tag.name,
                        "source_url": doc.metadata.get("url", ""),
                    })
                    doc.content += text + "\n"

            for table in soup.find_all("table"):
                table_data = []
                for row in table.find_all("tr"):
                    row_data = [cell.get_text().strip() for cell in row.find_all(["td", "th"])]
                    if row_data:
                        table_data.append(row_data)

                if table_data:
                    doc.tables.append({"data": table_data})

        except ImportError:
            self.logger.error("BeautifulSoup not available")
        except Exception as e:
            self.logger.error("html_parsing_failed", error=str(e))

        return doc

    async def _parse_text(self, file_path: str, file_name: str) -> ParsedDocument:
        """Parse plain text document"""
        doc = ParsedDocument()
        doc.metadata["file_name"] = file_name

        try:
            content = Path(file_path).read_text(encoding="utf-8")

            doc.title = file_name
            doc.content = content

            blocks = self._split_into_blocks(content)
            for block in blocks:
                doc.text_blocks.append({
                    "text": block,
                    "type": "text",
                })

        except Exception as e:
            self.logger.error("text_parsing_failed", error=str(e))

        return doc

    def _split_into_blocks(self, text: str, max_block_size: int = 1000) -> list[str]:
        """Split text into manageable blocks"""
        if not text:
            return []

        paragraphs = re.split(r"\n\s*\n", text)
        blocks = []

        for para in paragraphs:
            para = para.strip()
            if not para:
                continue

            if len(para) <= max_block_size:
                blocks.append(para)
            else:
                sentences = re.split(r"(?<=[。！？.!?])\s*", para)
                current = ""

                for sentence in sentences:
                    if len(current) + len(sentence) <= max_block_size:
                        current += sentence
                    else:
                        if current:
                            blocks.append(current.strip())
                        current = sentence

                if current.strip():
                    blocks.append(current.strip())

        return blocks

    def _extract_title(self, content: str) -> Optional[str]:
        """Extract title from document content"""
        lines = content.split("\n")

        for line in lines[:10]:
            line = line.strip()
            if len(line) > 5 and len(line) < 200:
                if line.startswith("#"):
                    return line.lstrip("#").strip()
                if line and line[0].isupper() or line[0] in "一二三四五":
                    return line

        return None
